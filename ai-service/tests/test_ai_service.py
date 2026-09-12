import os
from io import BytesIO

import httpx
import pytest
from docx import Document
from PIL import Image
from fastapi.testclient import TestClient
from tenacity import wait_none

from app.config import (
    EMBEDDING_DIMENSION,
    GEMINI_CHAT_COMPLETIONS_URL,
    GEMINI_DEFAULT_MODEL,
    GROQ_CHAT_COMPLETIONS_URL,
    GROQ_PRIMARY_MODEL,
    GROQ_SECONDARY_MODEL,
    HF_CHAT_COMPLETIONS_URL,
    HF_DEFAULT_MODEL,
    OPENROUTER_CHAT_COMPLETIONS_URL,
    OPENROUTER_DEFAULT_MODEL,
)
from app.main import app
from app.schemas import RagAnswerResponse, RagCitation
from app.services import agent as agent_service, extraction
from app.services.rag import HuggingFaceUpstreamError, build_grounded_prompt, call_hugging_face, create_citations

TOKEN = "tests-only-service-token"


class FakeEmbedding:
    def encode(self, texts, convert_to_numpy=True):
        return FakeEmbeddingArray([[float(index)] * EMBEDDING_DIMENSION for index, _ in enumerate(texts)])


class FakeEmbeddingArray:
    def __init__(self, values):
        self._values = values

    def tolist(self):
        return self._values


def client(rag_answer_generator=None) -> TestClient:
    os.environ["AI_SERVICE_TOKEN"] = TOKEN
    app.state.embedding_model = FakeEmbedding()
    if rag_answer_generator is None:
        if hasattr(app.state, "rag_answer_generator"):
            delattr(app.state, "rag_answer_generator")
    else:
        app.state.rag_answer_generator = rag_answer_generator

    return TestClient(app)


def headers() -> dict[str, str]:
    return {"X-Service-Token": TOKEN}


def test_extract_plain_text():
    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("note.txt", b"hello world", "text/plain")},
    )

    assert response.status_code == 200
    assert response.json()["text"] == "hello world"
    assert response.json()["contentType"] == "text/plain"


def test_extract_markdown():
    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("note.md", b"# Title\n\nBody", "text/markdown")},
    )

    assert response.status_code == 200
    assert "# Title" in response.json()["text"]


def test_extract_pdf():
    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("note.pdf", pdf_bytes(), "application/pdf")},
    )

    assert response.status_code == 200
    assert "Hello PDF" in response.json()["text"]


def test_text_pdf_does_not_trigger_scanned_pdf_ocr(monkeypatch):
    def unexpected_ocr(_data):
        raise AssertionError("OCR fallback must not run for normal text PDFs")

    monkeypatch.setattr(extraction, "extract_scanned_pdf_text", unexpected_ocr)

    assert "Hello PDF" in extraction.extract_pdf_text(pdf_bytes())


def test_scanned_pdf_heuristic_only_matches_an_absent_text_layer():
    assert extraction.needs_scanned_pdf_ocr([""]) is True
    assert extraction.needs_scanned_pdf_ocr(["Invoice 123"]) is False
    assert extraction.needs_scanned_pdf_ocr(["123"]) is False


def test_scanned_pdf_uses_ocr_fallback_when_embedded_text_is_absent(monkeypatch):
    class EmptyTextPage:
        def extract_text(self):
            return ""

    class ScannedReader:
        pages = [EmptyTextPage()]

    monkeypatch.setattr(extraction, "PdfReader", lambda _stream: ScannedReader())
    monkeypatch.setattr(extraction, "extract_scanned_pdf_text", lambda _data: "OCR recovered text")

    assert extraction.extract_pdf_text(b"scanned-pdf") == "OCR recovered text"


def test_extract_image_uses_ocr(monkeypatch):
    monkeypatch.setattr(extraction.pytesseract, "image_to_string", lambda _image: "OCR image text")
    image = Image.new("RGB", (8, 8), "white")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")

    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("scan.png", image_bytes.getvalue(), "image/png")},
    )

    assert response.status_code == 200
    assert response.json()["contentType"] == "image/png"
    assert response.json()["text"] == "OCR image text"


def test_extract_docx_includes_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Document heading")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    document_bytes = BytesIO()
    document.save(document_bytes)

    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("notes.docx", document_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    assert response.json()["text"] == "Document heading\nName\tValue"


def test_extract_csv_uses_existing_utf8_path():
    response = client().post(
        "/extract",
        headers=headers(),
        files={"file": ("items.csv", b"name,amount\nPens,3\n", "text/csv")},
    )

    assert response.status_code == 200
    assert response.json()["contentType"] == "text/csv"
    assert response.json()["text"] == "name,amount\nPens,3\n"


def test_extract_rejects_missing_service_token():
    response = client().post(
        "/extract",
        files={"file": ("note.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 401




def test_chunk_respects_word_boundaries_with_overlap():
    response = client().post(
        "/chunk",
        headers=headers(),
        json={"text": "alpha beta gamma delta", "chunkSize": 12, "overlap": 4},
    )

    assert response.status_code == 200
    assert response.json()["chunks"] == [
        {"index": 0, "text": "alpha beta ", "startOffset": 0, "endOffset": 11},
        {"index": 1, "text": "beta gamma ", "startOffset": 6, "endOffset": 17},
        {"index": 2, "text": "gamma delta", "startOffset": 11, "endOffset": 22},
    ]


def test_chunk_prefers_paragraph_boundaries():
    text = "Alpha paragraph.\n\nBeta paragraph.\n\nGamma paragraph."

    response = client().post(
        "/chunk",
        headers=headers(),
        json={"text": text, "chunkSize": 35, "overlap": 0},
    )

    assert response.status_code == 200
    assert response.json()["chunks"] == [
        {"index": 0, "text": text[:35], "startOffset": 0, "endOffset": 35},
        {"index": 1, "text": text[35:], "startOffset": 35, "endOffset": len(text)},
    ]


def test_chunk_falls_back_to_sentence_boundaries_for_long_paragraph():
    text = "First sentence. Second sentence. Third sentence."

    response = client().post(
        "/chunk",
        headers=headers(),
        json={"text": text, "chunkSize": 34, "overlap": 0},
    )

    assert response.status_code == 200
    assert response.json()["chunks"] == [
        {"index": 0, "text": text[:33], "startOffset": 0, "endOffset": 33},
        {"index": 1, "text": text[33:], "startOffset": 33, "endOffset": len(text)},
    ]


def test_chunk_keeps_single_long_word_intact():
    text = "x" * 2500

    response = client().post(
        "/chunk",
        headers=headers(),
        json={"text": text, "chunkSize": 1000, "overlap": 200},
    )

    assert response.status_code == 200
    assert response.json()["chunks"] == [
        {"index": 0, "text": text, "startOffset": 0, "endOffset": len(text)},
    ]


@pytest.mark.parametrize("text", ["", "short text"])
def test_chunk_handles_empty_and_short_text(text):
    response = client().post(
        "/chunk",
        headers=headers(),
        json={"text": text, "chunkSize": 1000, "overlap": 200},
    )

    assert response.status_code == 200
    expected = (
        []
        if not text
        else [{"index": 0, "text": text, "startOffset": 0, "endOffset": len(text)}]
    )
    assert response.json()["chunks"] == expected


def test_embed_returns_expected_shape():
    response = client().post(
        "/embed",
        headers=headers(),
        json={"texts": ["first", "second"]},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["dimension"] == EMBEDDING_DIMENSION
    assert len(body["embeddings"]) == 2
    assert len(body["embeddings"][0]) == EMBEDDING_DIMENSION


@pytest.mark.real_model
@pytest.mark.skipif(os.environ.get("RUN_REAL_EMBED_SMOKE") != "1", reason="set RUN_REAL_EMBED_SMOKE=1 to load the real embedding model")
def test_embed_real_model_smoke():
    os.environ["AI_SERVICE_TOKEN"] = TOKEN
    if hasattr(app.state, "embedding_model"):
        delattr(app.state, "embedding_model")

    try:
        response = TestClient(app).post(
            "/embed",
            headers=headers(),
            json={"texts": ["real model smoke test", "second embedding"]},
        )

        assert response.status_code == 200
        body = response.json()
        assert body["model"] == "sentence-transformers/all-MiniLM-L6-v2"
        assert body["dimension"] == EMBEDDING_DIMENSION
        assert len(body["embeddings"]) == 2
        assert len(body["embeddings"][0]) == EMBEDDING_DIMENSION
        assert all(isinstance(value, float) for value in body["embeddings"][0][:8])
    finally:
        if hasattr(app.state, "embedding_model"):
            delattr(app.state, "embedding_model")


def test_rag_prompt_includes_question_and_chunks():
    prompt = build_grounded_prompt(rag_request())

    assert "What is the renewal rule?" in prompt
    assert "[1] File: lease.pdf" in prompt
    assert "The lease renews yearly." in prompt
    assert "[2] File: policy.md" in prompt
    assert "Approval is required." in prompt
    assert "using only their numeric markers" in prompt


def test_rag_citations_include_only_valid_markers_in_first_mention_order():
    citations = create_citations(
        rag_request(),
        "Policy [2], repeated [2], then lease [1]. Invalid [99] and code `[1]` are ignored.",
    )

    assert [citation.reference for citation in citations] == ["2", "1"]
    assert [citation.fileName for citation in citations] == ["policy.md", "lease.pdf"]


def test_rag_answer_with_stubbed_llm_returns_expected_shape():
    def generator(payload):
        return RagAnswerResponse(
            answer="The lease renews yearly. [1]",
            model=HF_DEFAULT_MODEL,
            citations=[
                RagCitation(reference="1", fileName=payload.chunks[0].fileName, snippet=payload.chunks[0].text)
            ],
        )

    response = client(generator).post(
        "/rag/answer",
        headers=headers(),
        json=rag_request_json(),
    )

    assert response.status_code == 200
    body = response.json()
    assert body["answer"] == "The lease renews yearly. [1]"
    assert body["model"] == HF_DEFAULT_MODEL
    assert body["citations"] == [
        {"reference": "1", "fileName": "lease.pdf", "snippet": "The lease renews yearly."}
    ]


def test_rag_answer_with_upstream_failure_returns_clean_error():
    def generator(_payload):
        raise HuggingFaceUpstreamError("Hugging Face returned 503: overloaded")

    response = client(generator).post(
        "/rag/answer",
        headers=headers(),
        json=rag_request_json(),
    )

    assert response.status_code == 502
    assert response.json()["detail"] == "Hugging Face returned 503: overloaded"


def test_hugging_face_call_retries_transient_failures_then_succeeds(monkeypatch):
    monkeypatch.setenv("HF_API_TOKEN", "tests-only-hf-token")
    attempts = 0

    def fake_post(*_args, **_kwargs):
        nonlocal attempts
        attempts += 1
        if attempts <= 2:
            request = httpx.Request("POST", "https://huggingface.test/v1/chat/completions")
            raise httpx.ConnectError("temporary connection failure", request=request)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": "Recovered answer"}}],
                "usage": {"prompt_tokens": 21, "completion_tokens": 7},
            },
        )

    monkeypatch.setattr("app.services.rag.httpx.post", fake_post)

    result = call_hugging_face.retry_with(wait=wait_none())("grounded prompt", "test-model")

    assert result.answer == "Recovered answer"
    assert result.input_tokens == 21
    assert result.output_tokens == 7
    assert attempts == 3


def test_hugging_face_call_does_not_retry_non_transient_client_error(monkeypatch):
    monkeypatch.setenv("HF_API_TOKEN", "tests-only-hf-token")
    attempts = 0

    def fake_post(*_args, **_kwargs):
        nonlocal attempts
        attempts += 1
        return httpx.Response(400, text="invalid request")

    monkeypatch.setattr("app.services.rag.httpx.post", fake_post)

    with pytest.raises(HuggingFaceUpstreamError, match="returned 400"):
        call_hugging_face.retry_with(wait=wait_none())("grounded prompt", "test-model")

    assert attempts == 1


@pytest.mark.parametrize(
    ("environment", "failed_models", "expected_provider", "expected_model", "expected_url", "expected_calls"),
    [
        ({"GEMINI_API_KEY": "gemini-token"}, set(), "gemini", GEMINI_DEFAULT_MODEL, GEMINI_CHAT_COMPLETIONS_URL, 1),
        ({"GROQ_API_KEY": "groq-token"}, set(), "groq", GROQ_PRIMARY_MODEL, GROQ_CHAT_COMPLETIONS_URL, 1),
        (
            {"GROQ_API_KEY": "groq-token"},
            {GROQ_PRIMARY_MODEL},
            "groq",
            GROQ_SECONDARY_MODEL,
            GROQ_CHAT_COMPLETIONS_URL,
            2,
        ),
        (
            {"OPENROUTER_API_KEY": "openrouter-token"},
            set(),
            "openrouter",
            OPENROUTER_DEFAULT_MODEL,
            OPENROUTER_CHAT_COMPLETIONS_URL,
            1,
        ),
        ({"HF_API_TOKEN": "hf-token"}, set(), "huggingface", HF_DEFAULT_MODEL, HF_CHAT_COMPLETIONS_URL, 1),
    ],
)
def test_agent_completion_succeeds_for_each_provider_tier(
    monkeypatch,
    environment,
    failed_models,
    expected_provider,
    expected_model,
    expected_url,
    expected_calls,
):
    configure_agent_provider_environment(monkeypatch, environment)
    requests = []

    def fake_post(url, **kwargs):
        requests.append((url, kwargs))
        model = kwargs["json"]["model"]
        if model in failed_models:
            return httpx.Response(400, text="model rejected request")
        return provider_tool_call_response(expected_provider, model)

    monkeypatch.setattr("app.services.agent.httpx.post", fake_post)
    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 200
    assert len(requests) == expected_calls
    assert requests[-1][0] == expected_url
    expected_token = next(
        value for name, value in environment.items() if name.endswith("API_KEY") or name == "HF_API_TOKEN"
    )
    assert requests[-1][1]["headers"]["Authorization"] == f"Bearer {expected_token}"
    request_body = requests[-1][1]["json"]
    assert request_body["model"] == expected_model
    assert request_body["tool_choice"] == "auto"
    assert request_body["tools"][0]["function"]["name"] == "list_folder"
    body = response.json()
    assert body["model"] == expected_model
    assert body["usage"] == {"provider": expected_provider, "inputTokens": 12, "outputTokens": 4}
    assert body["message"]["toolCalls"][0]["function"] == {
        "name": "list_folder",
        "arguments": "{\"parentFolderId\":null}",
    }


def test_agent_completion_skips_unconfigured_tier_without_request(monkeypatch):
    configure_agent_provider_environment(monkeypatch, {"GROQ_API_KEY": "groq-token"})
    requests = []

    def fake_post(url, **kwargs):
        requests.append((url, kwargs["json"]["model"]))
        return provider_tool_call_response("groq", kwargs["json"]["model"])

    monkeypatch.setattr("app.services.agent.httpx.post", fake_post)
    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 200
    assert requests == [(GROQ_CHAT_COMPLETIONS_URL, GROQ_PRIMARY_MODEL)]


def test_agent_completion_uses_openrouter_model_override(monkeypatch):
    override_model = "vendor/rotating-tool-model:free"
    configure_agent_provider_environment(
        monkeypatch,
        {"OPENROUTER_API_KEY": "openrouter-token", "OPENROUTER_MODEL": override_model},
    )
    requested_models = []

    def fake_post(_url, **kwargs):
        requested_models.append(kwargs["json"]["model"])
        return provider_tool_call_response("openrouter", override_model)

    monkeypatch.setattr("app.services.agent.httpx.post", fake_post)
    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 200
    assert requested_models == [override_model]
    assert response.json()["model"] == override_model


def test_agent_completion_exhausts_transient_retries_then_falls_through(monkeypatch):
    configure_agent_provider_environment(
        monkeypatch,
        {"GEMINI_API_KEY": "gemini-token", "GROQ_API_KEY": "groq-token"},
    )
    attempts = []

    def fake_post(url, **kwargs):
        model = kwargs["json"]["model"]
        attempts.append(model)
        if model == GEMINI_DEFAULT_MODEL:
            request = httpx.Request("POST", url)
            raise httpx.ConnectError("temporary connection failure", request=request)
        return provider_tool_call_response("groq", model)

    no_wait_call = agent_service.call_openai_compatible_with_tools.retry_with(wait=wait_none())
    monkeypatch.setattr(agent_service, "call_openai_compatible_with_tools", no_wait_call)
    monkeypatch.setattr("app.services.agent.httpx.post", fake_post)
    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 200
    assert attempts == [GEMINI_DEFAULT_MODEL, GEMINI_DEFAULT_MODEL, GROQ_PRIMARY_MODEL]
    assert response.json()["usage"]["provider"] == "groq"


def test_agent_completion_reports_total_chain_exhaustion(monkeypatch):
    configure_agent_provider_environment(
        monkeypatch,
        {
            "GEMINI_API_KEY": "gemini-token",
            "GROQ_API_KEY": "groq-token",
            "OPENROUTER_API_KEY": "openrouter-token",
            "HF_API_TOKEN": "hf-token",
        },
    )
    attempted_models = []

    def fake_post(_url, **kwargs):
        attempted_models.append(kwargs["json"]["model"])
        return httpx.Response(400, text="invalid request")

    monkeypatch.setattr("app.services.agent.httpx.post", fake_post)
    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 502
    assert attempted_models == [
        GEMINI_DEFAULT_MODEL,
        GROQ_PRIMARY_MODEL,
        GROQ_SECONDARY_MODEL,
        OPENROUTER_DEFAULT_MODEL,
        HF_DEFAULT_MODEL,
    ]
    detail = response.json()["detail"]
    assert detail.startswith("All configured agent completion providers failed:")
    for provider_name, model in (
        ("gemini", GEMINI_DEFAULT_MODEL),
        ("groq", GROQ_PRIMARY_MODEL),
        ("groq", GROQ_SECONDARY_MODEL),
        ("openrouter", OPENROUTER_DEFAULT_MODEL),
        ("huggingface", HF_DEFAULT_MODEL),
    ):
        assert f"{provider_name}/{model}" in detail


def test_agent_completion_reports_when_no_provider_is_configured(monkeypatch):
    configure_agent_provider_environment(monkeypatch, {})
    monkeypatch.setattr(
        "app.services.agent.httpx.post",
        lambda *_args, **_kwargs: pytest.fail("No HTTP request should be attempted"),
    )

    response = client().post("/agent/complete", headers=headers(), json=agent_request_json())

    assert response.status_code == 503
    assert response.json()["detail"] == "No agent completion provider API keys are configured"


def configure_agent_provider_environment(monkeypatch, values):
    for name in ("GEMINI_API_KEY", "GROQ_API_KEY", "OPENROUTER_API_KEY", "OPENROUTER_MODEL", "HF_API_TOKEN", "HF_MODEL"):
        monkeypatch.delenv(name, raising=False)
    for name, value in values.items():
        monkeypatch.setenv(name, value)


def agent_request_json() -> dict:
    return {
        "messages": [
            {"role": "system", "content": "Use tools safely."},
            {"role": "user", "content": "List root."},
        ],
        "tools": [
            {
                "type": "function",
                "function": {
                    "name": "list_folder",
                    "description": "List files.",
                    "parameters": {"type": "object"},
                },
            }
        ],
    }


def provider_tool_call_response(provider: str, model: str) -> httpx.Response:
    body = {
        "id": f"chatcmpl-{provider}",
        "object": "chat.completion",
        "model": model,
        "choices": [
            {
                "index": 0,
                "message": {
                    "role": "assistant",
                    "content": None,
                    "tool_calls": [
                        {
                            "id": f"call-{provider}",
                            "type": "function",
                            "function": {
                                "name": "list_folder",
                                "arguments": "{\"parentFolderId\":null}",
                            },
                        }
                    ],
                },
                "finish_reason": "tool_calls",
            }
        ],
        "usage": {"prompt_tokens": 12, "completion_tokens": 4, "total_tokens": 16},
    }
    if provider == "groq":
        body["x_groq"] = {"id": "req-groq", "model_version": model}
    return httpx.Response(200, json=body)


def rag_request_json() -> dict:
    return {
        "question": "What is the renewal rule?",
        "chunks": [
            {"fileName": "lease.pdf", "text": "The lease renews yearly."},
            {"fileName": "policy.md", "text": "Approval is required."},
        ],
        "history": [
            {"role": "user", "content": "What document is this?"},
            {"role": "assistant", "content": "It is a lease."},
        ],
    }


def rag_request():
    from app.schemas import RagAnswerRequest

    return RagAnswerRequest(**rag_request_json())


def pdf_bytes() -> bytes:
    return b"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> /MediaBox [0 0 612 792] /Contents 5 0 R >>
endobj
4 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
5 0 obj
<< /Length 44 >>
stream
BT /F1 24 Tf 100 700 Td (Hello PDF) Tj ET
endstream
endobj
xref
0 6
0000000000 65535 f
0000000009 00000 n
0000000058 00000 n
0000000115 00000 n
0000000241 00000 n
0000000311 00000 n
trailer
<< /Root 1 0 R /Size 6 >>
startxref
405
%%EOF
"""
